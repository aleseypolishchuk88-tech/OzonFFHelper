"""Тесты парсера описи на обезличенном файле tests/fixtures/Опись_sample.xlsx."""
from datetime import date
from pathlib import Path

import pytest

from core.opis_parser import OpisParseError, parse_opis

FIXTURE = Path(__file__).parent / "fixtures" / "Опись_sample.xlsx"


@pytest.fixture(scope="module")
def opis():
    return parse_opis(FIXTURE)


def test_header_fields(opis):
    assert opis.opis_number == "10001"
    assert opis.opis_date == date(2026, 5, 15)
    assert opis.supply_order_number == "2000000000001"
    assert "ИВАНОВО" in opis.warehouse_ozon.upper()
    assert opis.counterparty  # ФИО заполнено


def test_three_boxes(opis):
    """Ключевое требование брифа: ровно 3 короба."""
    assert opis.total_boxes == 3
    assert [b.pac_code for b in opis.boxes] == [
        "PAC100000001",
        "PAC100000002",
        "PAC100000003",
    ]


def test_products_and_quantities(opis):
    skus = [p.ozn_sku for b in opis.boxes for p in b.products]
    assert skus == ["OZN100000001", "OZN100000002", "OZN100000003"]
    # по 10 шт в каждом коробе, всего 30
    assert all(b.total_quantity == 10 for b in opis.boxes)
    assert opis.total_quantity == 30


def test_box_metrics(opis):
    first = opis.boxes[0]
    assert first.weight_kg == pytest.approx(2.9)
    assert first.volume_m3 == pytest.approx(0.038)


def test_ignores_itogo_row(opis):
    # "Итого" не должно стать коробом/товаром
    assert all("ИТОГ" not in b.pac_code.upper() for b in opis.boxes)


def test_bad_file_raises(tmp_path):
    empty = tmp_path / "empty.xlsx"
    import openpyxl

    openpyxl.Workbook().save(empty)
    with pytest.raises(OpisParseError):
        parse_opis(empty)


def _fixture_copy(tmp_path, edits):
    import openpyxl

    wb = openpyxl.load_workbook(FIXTURE)
    ws = wb.active
    for coord, val in edits.items():
        ws[coord].value = val
    path = tmp_path / "opis.xlsx"
    wb.save(path)
    return path


def test_duplicate_pac_raises(tmp_path):
    # второй короб назван так же, как первый
    path = _fixture_copy(tmp_path, {"A15": "PAC100000001"})
    with pytest.raises(OpisParseError, match="дважды"):
        parse_opis(path)


def test_unrecognized_code_with_quantity_raises(tmp_path):
    # опечатка в коде товара: товар не должен молча выпасть
    path = _fixture_copy(tmp_path, {"A16": "0ZN100000002"})
    with pytest.raises(OpisParseError, match="0ZN100000002"):
        parse_opis(path)
